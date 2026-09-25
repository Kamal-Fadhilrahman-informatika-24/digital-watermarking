"""Pengujian alur web end-to-end lewat Flask test client (embed -> detect -> attack -> comparison)."""

import io
import re

import cv2
import numpy as np
import pytest
from openpyxl import load_workbook

from tests.conftest import upload


@pytest.mark.parametrize("url", ["/", "/embed", "/detect", "/testing", "/comparison", "/about"])
def test_main_pages_render(client, url):
    response = client.get(url)
    assert response.status_code == 200
    assert b"<html" in response.data and len(response.data) > 1500


def _embed(client, cover_png, key, method="dct", text="KAMAL FADHILRAHMAN"):
    response = upload(client, "/embed", cover_png, "foto saya.png", watermark_text=text, secret_key=key, method=method)
    assert response.status_code == 302, response.get_data(as_text=True)[:500]
    return response.headers["Location"].rsplit("/", 1)[1]


def test_embed_result_and_download(client, cover_png, key):
    job_id = _embed(client, cover_png, key)
    page = client.get(f"/result/{job_id}").get_data(as_text=True)
    assert "PSNR" in page and "KAMAL FADHILRAHMAN" in page and "dB" in page
    download = client.get(f"/download/{job_id}")
    assert download.status_code == 200 and download.data[:8] == b"\x89PNG\r\n\x1a\n"
    assert "attachment" in download.headers["Content-Disposition"]
    served = client.get(f"/files/watermarked/{job_id}.png")
    assert served.status_code == 200 and served.mimetype == "image/png"
    assert client.get(f"/files/original/{job_id}.png").status_code == 200


def test_watermarked_download_is_really_watermarked(client, cover_png, key, other_key):
    job_id = _embed(client, cover_png, key)
    data = client.get(f"/download/{job_id}").data
    ok = client.post("/detect", data={"image": (io.BytesIO(data), "wm.png"), "secret_key": key, "method": "dct",
                                       "expected": "KAMAL FADHILRAHMAN"}, content_type="multipart/form-data")
    assert ok.status_code == 200 and b"SUCCESS" in ok.data and "KAMAL FADHILRAHMAN" in ok.get_data(as_text=True)
    wrong = client.post("/detect", data={"image": (io.BytesIO(data), "wm.png"), "secret_key": other_key, "method": "dct"},
                        content_type="multipart/form-data")
    assert wrong.status_code == 200 and b"FAIL" in wrong.data and b"SUCCESS" not in wrong.data


def test_detect_reuses_previous_job_without_upload(client, cover_png, key):
    job_id = _embed(client, cover_png, key)
    response = client.post("/detect", data={"job": job_id, "secret_key": key})
    assert response.status_code == 200 and b"SUCCESS" in response.data


def test_detect_rejects_bad_original_size(client, cover_png, key):
    job_id = _embed(client, cover_png, key)
    response = client.post("/detect", data={"job": job_id, "secret_key": key, "orig_w": "abc", "orig_h": "10"})
    assert response.status_code == 400 and "Ukuran asli" in response.get_data(as_text=True)


def test_detect_after_resize_with_original_size(client, cover, cover_png, key):
    job_id = _embed(client, cover_png, key)
    wm = cv2.imdecode(np.frombuffer(client.get(f"/download/{job_id}").data, np.uint8), cv2.IMREAD_COLOR)
    small = cv2.resize(wm, (384, 384), interpolation=cv2.INTER_AREA)
    data = cv2.imencode(".png", small)[1].tobytes()
    response = client.post("/detect", data={"image": (io.BytesIO(data), "kecil.png"), "secret_key": key, "method": "dct",
                                             "orig_w": "512", "orig_h": "512"}, content_type="multipart/form-data")
    assert response.status_code == 200 and b"SUCCESS" in response.data


def test_lsb_method_via_web(client, cover_png, key):
    job_id = _embed(client, cover_png, key, method="lsb")
    page = client.get(f"/result/{job_id}").get_data(as_text=True)
    assert "LSB Fragile" in page
    assert b"SUCCESS" in client.post("/detect", data={"job": job_id, "secret_key": key}).data


def test_large_image_is_downscaled_before_processing(client, key):
    rng = np.random.default_rng(0)
    big = cv2.GaussianBlur(rng.integers(0, 256, (1400, 1800, 3), dtype=np.uint8), (0, 0), 3)
    job_id = _embed(client, cv2.imencode(".png", big)[1].tobytes(), key)
    page = client.get(f"/result/{job_id}").get_data(as_text=True)
    assert "1024" in page  # sisi terpanjang dibatasi MAX_PROCESS_SIDE


def test_attack_testing_flow_with_nested_attack_files(client, storage, cover_png, key):
    job_id = _embed(client, cover_png, key)
    response = client.post("/testing", data={"job": job_id, "secret_key": key})
    assert response.status_code == 302 and "/testing/run/" in response.headers["Location"]
    run_id = response.headers["Location"].rsplit("/", 1)[1]

    page = client.get(f"/testing/run/{run_id}").get_data(as_text=True)
    for label in ("JPEG", "Cropping", "Resize", "Gaussian Noise", "Brightness", "Contrast", "SUCCESS"):
        assert label in page, label
    assert "attack_vs_psnr.png" in page

    # file serangan berada di subfolder run_id -> path bersarang harus terlayani
    attack_urls = re.findall(r'/files/attacks/[0-9a-f]{12}/[\w\-.]+\.(?:png|jpg)', page)
    assert len(set(attack_urls)) == 9
    for url in set(attack_urls):
        r = client.get(url)
        assert r.status_code == 200 and len(r.data) > 1000, url

    chart = re.search(r'/files/charts/[\w\-.]+\.png', page).group(0)
    assert client.get(chart).status_code == 200

    xlsx = client.get(f"/report/run_{run_id}.xlsx")
    assert xlsx.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx.data))
    assert "Attack Results" in workbook.sheetnames
    assert workbook["Attack Results"].max_row - 1 == 9


def test_attack_testing_with_wrong_key_reports_failure(client, cover_png, key, other_key):
    job_id = _embed(client, cover_png, key)
    response = client.post("/testing", data={"job": job_id, "secret_key": other_key})
    page = client.get(response.headers["Location"]).get_data(as_text=True)
    assert "secret key salah" in page  # peringatan key ditolak
    assert 'class="badge ok"' not in page and page.count('class="badge fail"') == 9


def test_attack_testing_with_fresh_upload(client, cover_png, key):
    response = upload(client, "/testing", cover_png, watermark_text="UJI", secret_key=key, method="dct")
    assert response.status_code == 302 and "/testing/run/" in response.headers["Location"]


def test_attack_testing_requires_key(client, cover_png, key):
    job_id = _embed(client, cover_png, key)
    response = client.post("/testing", data={"job": job_id, "secret_key": ""})
    assert response.status_code == 400


def test_comparison_page_empty_then_filled(app, client, tmp_path, key):
    empty = client.get("/comparison")
    assert empty.status_code == 200 and b"Belum ada hasil" in empty.data
    response = client.post("/comparison/run", data={"secret_key": key, "watermark_text": "UJI"})
    assert response.status_code == 302  # dataset kosong -> dibuat otomatis lalu dievaluasi
    page = client.get("/comparison").get_data(as_text=True)
    assert "DCT Robust" in page and "LSB Fragile" in page and "JPEG 90" in page
    assert client.get("/report/watermark_testing.xlsx").status_code == 200


def test_comparison_run_requires_key(client):
    response = client.post("/comparison/run", data={"secret_key": "", "watermark_text": "UJI"})
    assert response.status_code == 400
