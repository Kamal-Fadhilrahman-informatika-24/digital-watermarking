"""Pengujian laporan: ringkasan, grafik PNG, XLSX, dan simpan/muat JSON hasil evaluasi."""

import math

import pytest
from openpyxl import load_workbook

from watermark.evaluation import (
    evaluate_dataset,
    generate_synthetic_dataset,
    list_dataset_images,
    load_results_json,
    save_results_json,
)
from watermark.report import comparison_facts, make_charts, summarize, write_xlsx


@pytest.fixture(scope="module")
def results(tmp_path_factory, key, text):
    base = tmp_path_factory.mktemp("eval")
    dataset = base / "original"
    generate_synthetic_dataset(dataset, size=512)
    for extra in sorted(dataset.iterdir())[2:]:  # 2 citra cukup untuk pengujian (lebih cepat)
        extra.unlink()
    return evaluate_dataset(dataset, key, text, watermarked_dir=base / "wm", attack_dir=base / "atk")


def test_generate_dataset_creates_named_images(tmp_path):
    written = generate_synthetic_dataset(tmp_path, size=256)
    assert [p.name for p in written] == [f"image0{i}.png" for i in range(1, 6)]
    assert len(list_dataset_images(tmp_path)) == 5
    assert generate_synthetic_dataset(tmp_path, size=256) == [], "tidak menimpa file yang sudah ada"


def test_evaluation_covers_every_image_method_and_attack(results, key):
    assert len(results["dataset"]) == 2
    assert len(results["rows"]) == 2 * 2 * 9
    assert results["methods"] == ["DCT Robust", "LSB Fragile"]
    assert key not in str(results), "secret key tidak boleh ikut tersimpan di hasil evaluasi"


def test_summary_counts_are_consistent(results):
    summary = summarize(results["rows"])
    assert len(summary) == 2 * 9
    for item in summary:
        assert item["n"] == 2 and 0 <= item["success"] <= 2
        assert 0.0 <= item["mean_ber"] <= 1.0
    dct_jpeg = next(s for s in summary if s["method_key"] == "dct" and s["attack_key"] == "jpeg_50")
    lsb_jpeg = next(s for s in summary if s["method_key"] == "lsb" and s["attack_key"] == "jpeg_50")
    assert dct_jpeg["mean_ber"] < lsb_jpeg["mean_ber"]


def test_comparison_facts_are_derived_from_data(results):
    facts = comparison_facts(results["rows"])
    assert len(facts) == 9
    assert all("DCT Robust" in f and "LSB Fragile" in f for f in facts)


def test_json_roundtrip_preserves_infinite_psnr(results, tmp_path):
    results["rows"][0]["psnr_after_attack"] = math.inf  # kasus MSE = 0
    path = tmp_path / "hasil.json"
    save_results_json(results, path)
    assert "Infinity" not in path.read_text(), "JSON harus valid (tanpa token Infinity)"
    loaded = load_results_json(path)
    assert loaded["rows"][0]["psnr_after_attack"] == math.inf
    assert load_results_json(tmp_path / "tidak_ada.json") is None


def test_charts_are_real_png_files(results, tmp_path):
    files = make_charts(results["rows"], tmp_path, prefix="uji_")
    assert set(files) == {"psnr", "nc", "ber", "success"}
    for name in files.values():
        data = (tmp_path / name).read_bytes()
        assert data[:8] == b"\x89PNG\r\n\x1a\n" and len(data) > 5000


def test_xlsx_structure_and_content(results, tmp_path, key):
    path = write_xlsx(results, tmp_path / "laporan.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Summary", "DCT Results", "LSB Results", "Attack Results", "Dataset"]
    ws = workbook["Attack Results"]
    header = [c.value for c in ws[1]]
    assert header[:2] == ["Image", "Method"] and "Status" in header and "BER" in header
    assert ws.max_row - 1 == len(results["rows"])
    statuses = {row[header.index("Status")].value for row in ws.iter_rows(min_row=2)}
    assert statuses <= {"SUCCESS", "FAIL"}
    assert workbook["Dataset"].max_row - 1 == 2


def test_xlsx_never_contains_secret_key(results, tmp_path, key):
    workbook = load_workbook(write_xlsx(results, tmp_path / "laporan.xlsx"))
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            assert not any(isinstance(v, str) and key in v for v in row)
