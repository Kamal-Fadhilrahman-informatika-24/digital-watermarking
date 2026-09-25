"""Pengujian keamanan: secret key, upload, path traversal, error page, kebocoran data."""

import json
import logging
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.conftest import upload  # noqa: F401  (helper multipart)
from watermark import config
from watermark.key_utils import (
    derive_seed,
    generate_block_positions,
    generate_lsb_positions,
    generate_pair_indices,
)
from watermark.utils import allowed_file, cleanup_old_files, is_valid_job_id

ROOT = Path(__file__).resolve().parent.parent


# ----------------------------------------------------------- secret key
def test_key_determinism():
    assert derive_seed("abc") == derive_seed("abc")
    assert derive_seed("abc") != derive_seed("abd")
    assert generate_block_positions(500, "abc") == generate_block_positions(500, "abc")
    assert generate_block_positions(500, "abc") != generate_block_positions(500, "abd")
    assert (generate_pair_indices(500, 3, "abc") == generate_pair_indices(500, 3, "abc")).all()
    assert generate_lsb_positions(10_000, 100, "abc") == generate_lsb_positions(10_000, 100, "abc")


def test_block_order_is_a_permutation():
    order = generate_block_positions(1000, "kunci")
    assert sorted(order) == list(range(1000))
    assert order != list(range(1000))


def test_lsb_positions_are_distinct_and_prefix_consistent():
    long = generate_lsb_positions(5000, 300, "kunci")
    assert len(set(long)) == 300
    assert generate_lsb_positions(5000, 48, "kunci") == long[:48]


def test_contexts_give_independent_streams():
    assert derive_seed("kunci", "dct-block-order") != derive_seed("kunci", "lsb-positions")


def test_empty_key_is_rejected():
    with pytest.raises(ValueError):
        derive_seed("")


def _source_files():
    files = [ROOT / "app.py"]
    for folder in ("watermark", "routes", "scripts"):
        files += sorted((ROOT / folder).glob("*.py"))
    files += sorted((ROOT / "templates").glob("*.html"))
    return files


def test_no_hardcoded_secrets_in_source():
    patterns = [
        re.compile(r"""SECRET_KEY["']?\]?\s*=\s*["'][^"']+["']"""),
        re.compile(r"""(password|passwd|api[_-]?key|token)\s*=\s*["'][^"']{6,}["']""", re.I),
        re.compile(r"BEGIN (RSA |EC )?PRIVATE KEY"),
        re.compile(r"[0-9a-f]{40,}", re.I),
    ]
    offenders = []
    for path in _source_files():
        text = path.read_text(encoding="utf-8")
        for index, pattern in enumerate(patterns):
            if index == 1 and path.suffix != ".py":  # atribut HTML seperti data-toggle-password="..." bukan kredensial
                continue
            for match in pattern.finditer(text):
                offenders.append(f"{path.name}: {match.group(0)[:40]}")
    assert not offenders, offenders


def test_flask_secret_key_comes_from_environment_or_random():
    source = (ROOT / "app.py").read_text(encoding="utf-8")
    assert 'os.environ.get("FLASK_SECRET_KEY")' in source and "secrets.token_hex" in source


def test_env_example_contains_no_real_secret():
    content = (ROOT / ".env.example").read_text(encoding="utf-8")
    assert re.search(r"^FLASK_SECRET_KEY=\s*$", content, re.M)


def test_gitignore_protects_sensitive_and_generated_files():
    lines = (ROOT / ".gitignore").read_text(encoding="utf-8").splitlines()
    for required in (".venv/", "__pycache__/", ".DS_Store", ".env", "uploads/*", "outputs/*"):
        assert required in lines, f".gitignore belum memuat {required}"


def test_no_env_file_is_committed_to_project():
    assert not (ROOT / ".env").exists(), "jangan menyimpan .env di dalam folder proyek yang di-commit"


# --------------------------------------------------------------- upload
def test_upload_rejects_wrong_extension(client, cover_png, key):
    response = upload(client, "/embed", cover_png, "berkas.txt", watermark_text="ABC", secret_key=key, method="dct")
    assert response.status_code == 400 and b"Format file tidak didukung" in response.data


def test_upload_rejects_fake_png(client, key):
    response = upload(client, "/embed", b"ini bukan gambar sama sekali", "palsu.png", watermark_text="ABC", secret_key=key, method="dct")
    assert response.status_code == 400 and b"bukan gambar" in response.data


def test_upload_rejects_corrupt_png(client, key):
    corrupt = b"\x89PNG\r\n\x1a\n" + b"\x00" * 50
    response = upload(client, "/embed", corrupt, "rusak.png", watermark_text="ABC", secret_key=key, method="dct")
    assert response.status_code == 400 and b"Traceback" not in response.data


def test_upload_rejects_empty_file(client, key):
    response = upload(client, "/embed", b"", "kosong.png", watermark_text="ABC", secret_key=key, method="dct")
    assert response.status_code == 400


def test_upload_without_file_is_friendly(client, key):
    response = client.post("/embed", data={"watermark_text": "ABC", "secret_key": key, "method": "dct"})
    assert response.status_code == 400 and b"Pilih file gambar" in response.data


def test_upload_rejects_too_small_image(client, key):
    import cv2
    import numpy as np

    tiny = cv2.imencode(".png", np.zeros((64, 64, 3), np.uint8))[1].tobytes()
    response = upload(client, "/embed", tiny, "kecil.png", watermark_text="ABC", secret_key=key, method="dct")
    assert response.status_code == 400 and b"terlalu kecil" in response.data


def test_upload_too_large_returns_413(app, client, cover_png, key):
    app.config["MAX_CONTENT_LENGTH"] = 1000
    response = upload(client, "/embed", cover_png, "besar.png", watermark_text="ABC", secret_key=key, method="dct")
    assert response.status_code == 413 and b"terlalu besar" in response.data


def test_missing_secret_key_and_text_are_rejected(client, cover_png, key):
    no_key = upload(client, "/embed", cover_png, watermark_text="ABC", secret_key="", method="dct")
    no_text = upload(client, "/embed", cover_png, watermark_text="", secret_key=key, method="dct")
    bad_method = upload(client, "/embed", cover_png, watermark_text="ABC", secret_key=key, method="rot13")
    too_long = upload(client, "/embed", cover_png, watermark_text="x" * 200, secret_key=key, method="dct")
    for response in (no_key, no_text, bad_method, too_long):
        assert response.status_code == 400 and b"Traceback" not in response.data


def test_allowed_file_extensions():
    assert allowed_file("a.PNG") and allowed_file("a.jpeg") and allowed_file("a.b.jpg")
    assert not allowed_file("a.exe") and not allowed_file("a.png.php") and not allowed_file("png")


def test_uploaded_filename_is_sanitised(client, storage, cover_png, key):
    response = upload(client, "/embed", cover_png, "../../etc/evil name.png", watermark_text="ABC", secret_key=key, method="dct")
    assert response.status_code == 302
    stored = [p.name for p in storage.upload.iterdir() if not p.name.startswith(".")]
    assert len(stored) == 1 and re.fullmatch(r"[0-9a-f]{12}\.png", stored[0])
    job = json.loads(next(storage.watermarked.glob("*.json")).read_text(encoding="utf-8"))
    assert "/" not in job["filename"] and ".." not in job["filename"]
    assert not (storage.upload.parent / "etc").exists()


# ------------------------------------------------------- path traversal
@pytest.mark.parametrize("url", [
    "/files/watermarked/../../app.py",
    "/files/watermarked/%2e%2e/%2e%2e/app.py",
    "/files/watermarked/..%2f..%2fapp.py.png",
    "/files/attacks/../reports/run_000000000000.json",
    "/files/secret/anything.png",
    "/files/watermarked/data.json",
    "/files/watermarked/x.png.txt",
    "/result/..%2f..%2fapp.py",
    "/result/notavalidid",
    "/download/../app.py",
    "/download/zzzzzzzzzzzz",
    "/report/../app.py",
    "/report/..%2fapp.py",
    "/report/run_000000000000.xlsx",
    "/report/last_evaluation.json",
    "/testing/run/../../app.py",
    "/testing/run/000000000000",
])
def test_path_traversal_and_unknown_files_return_404(client, url):
    response = client.get(url)
    assert response.status_code == 404
    assert b"import " not in response.data and b"SECRET" not in response.data


def test_job_id_validation():
    assert is_valid_job_id("0123456789ab")
    for bad in ("", "../x", "0123456789AB", "0123456789abc", "0123456789a", "..", None):
        assert not is_valid_job_id(bad)


# ------------------------------------------------------------ error page
def test_unexpected_error_does_not_leak_details(client, cover_png, key, monkeypatch):
    def boom(*args, **kwargs):
        raise RuntimeError("DETAIL-INTERNAL /home/claude/rahasia")

    monkeypatch.setattr("routes.watermark_routes.embed_upload", boom)
    response = upload(client, "/embed", cover_png, watermark_text="ABC", secret_key=key, method="dct")
    assert response.status_code == 500
    body = response.get_data(as_text=True)
    assert "Terjadi kesalahan" in body
    assert "DETAIL-INTERNAL" not in body and "Traceback" not in body and "/home/claude" not in body


def test_404_and_405_pages_are_friendly(client):
    assert b"tidak ditemukan" in client.get("/halaman-tidak-ada").data
    response = client.post("/about")
    assert response.status_code == 405 and b"Traceback" not in response.data


# ------------------------------------------------------- kebocoran secret
def test_secret_key_is_never_stored_logged_or_rendered(client, storage, cover_png, caplog):
    secret = "RAHASIA-uji-jangan-bocor-9271"
    caplog.set_level(logging.DEBUG)
    response = upload(client, "/embed", cover_png, watermark_text="KAMAL", secret_key=secret, method="dct")
    assert response.status_code == 302
    job_id = response.headers["Location"].rsplit("/", 1)[1]
    page = client.get(f"/result/{job_id}").get_data(as_text=True)
    assert secret not in page

    run_response = client.post("/testing", data={"job": job_id, "secret_key": secret})
    assert run_response.status_code == 302
    run_page = client.get(run_response.headers["Location"]).get_data(as_text=True)
    assert secret not in run_page

    assert secret not in caplog.text, "secret key tidak boleh muncul di log"
    for path in storage.upload.parent.rglob("*"):
        if path.is_file() and path.suffix in {".json", ".txt"}:
            assert secret not in path.read_text(encoding="utf-8", errors="ignore")
    for xlsx in storage.reports.glob("*.xlsx"):
        for sheet in load_workbook(xlsx).worksheets:
            for row in sheet.iter_rows(values_only=True):
                assert not any(isinstance(v, str) and secret in v for v in row)


def test_secret_input_is_a_password_field(client):
    for url in ("/embed", "/detect", "/testing"):
        html = client.get(url).get_data(as_text=True)
        assert re.search(r'<input[^>]*type="password"[^>]*name="secret_key"|<input[^>]*name="secret_key"[^>]*type="password"', html), url
        assert 'autocomplete="off"' in html or "autocomplete='off'" in html or 'autocomplete="new-password"' in html, url


# ----------------------------------------------------------- housekeeping
def test_cleanup_removes_only_old_files_and_keeps_gitkeep(tmp_path):
    import os
    import time

    old, new, keep = tmp_path / "lama.png", tmp_path / "baru.png", tmp_path / ".gitkeep"
    for f in (old, new, keep):
        f.write_bytes(b"x")
    long_ago = time.time() - 3 * 24 * 3600
    os.utime(old, (long_ago, long_ago))
    os.utime(keep, (long_ago, long_ago))
    assert cleanup_old_files(tmp_path, max_age_hours=24) == 1
    assert not old.exists() and new.exists() and keep.exists()


def test_upload_limit_is_configured():
    assert config.MAX_UPLOAD_BYTES == 8 * 1024 * 1024
    assert config.ALLOWED_EXTENSIONS == {"png", "jpg", "jpeg"}
