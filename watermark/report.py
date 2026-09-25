"""Report generator: summary tables, Matplotlib charts and the XLSX workbook."""

from __future__ import annotations

import math
from pathlib import Path
from typing import Dict, List

import matplotlib

matplotlib.use("Agg")  # tanpa GUI; aman dijalankan di server Flask
import matplotlib.pyplot as plt  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

# ---------------------------------------------------------------- summary
def _finite_mean(values: List[float]) -> float:
    finite = [v for v in values if isinstance(v, (int, float)) and math.isfinite(v)]
    if finite:
        return sum(finite) / len(finite)
    return math.inf if values else math.nan


def summarize(rows: List[Dict]) -> List[Dict]:
    """Aggregate rows per (attack, method): mean PSNR/NC/BER and success count."""
    order: List[tuple] = []
    groups: Dict[tuple, List[Dict]] = {}
    for row in rows:
        key = (row["attack_key"], row["method_key"])
        if key not in groups:
            order.append(key)
            groups[key] = []
        groups[key].append(row)
    summary = []
    for key in order:
        items = groups[key]
        first = items[0]
        summary.append({
            "attack_key": key[0], "method_key": key[1], "method": first["method"],
            "attack": first["attack"], "parameter": first["parameter"], "short": first.get("short", first["attack"]),
            "n": len(items),
            "mean_psnr": _finite_mean([r["psnr_after_attack"] for r in items]),
            "mean_nc": sum(r["nc"] for r in items) / len(items),
            "mean_ber": sum(r["ber"] for r in items) / len(items),
            "success": sum(1 for r in items if r["status"] == "SUCCESS"),
        })
    return summary


def comparison_facts(rows: List[Dict]) -> List[str]:
    """Factual, data-derived sentences comparing methods per attack (no claims beyond the data)."""
    summary = summarize(rows)
    attacks: List[str] = []
    for s in summary:
        if s["attack_key"] not in attacks:
            attacks.append(s["attack_key"])
    facts = []
    for attack_key in attacks:
        parts = [s for s in summary if s["attack_key"] == attack_key]
        if len(parts) < 2:
            continue
        desc = ", ".join(f"{p['method']}: {p['success']}/{p['n']} berhasil, BER {p['mean_ber']:.4f}, NC {p['mean_nc']:.4f}" for p in parts)
        best = min(parts, key=lambda p: p["mean_ber"])
        tie = all(abs(p["mean_ber"] - best["mean_ber"]) < 1e-12 for p in parts)
        verdict = "BER sama pada kedua metode." if tie else f"BER terendah: {best['method']}."
        facts.append(f"{parts[0]['short']} - {desc}. {verdict}")
    return facts


# ----------------------------------------------------------------- charts
def _grouped_bar(summary: List[Dict], value_key: str, ylabel: str, title: str, path: Path, ylim=None) -> None:
    methods: List[str] = []
    attacks: List[tuple] = []
    for s in summary:
        if s["method"] not in methods:
            methods.append(s["method"])
        if (s["attack_key"], s["short"]) not in attacks:
            attacks.append((s["attack_key"], s["short"]))
    fig, ax = plt.subplots(figsize=(10, 4.8))
    width = 0.8 / max(1, len(methods))
    for m_index, method in enumerate(methods):
        values = []
        for attack_key, _ in attacks:
            match = [s for s in summary if s["method"] == method and s["attack_key"] == attack_key]
            value = match[0][value_key] if match else math.nan
            values.append(value if math.isfinite(value) else math.nan)
        positions = [i + m_index * width - 0.4 + width / 2 for i in range(len(attacks))]
        bars = ax.bar(positions, values, width, label=method)
        for bar, value in zip(bars, values):
            if math.isfinite(value):
                ax.annotate(f"{value:.2f}" if value >= 10 else f"{value:.3f}", (bar.get_x() + bar.get_width() / 2, value),
                            ha="center", va="bottom", fontsize=6.5, xytext=(0, 1), textcoords="offset points")
    ax.set_xticks(range(len(attacks)))
    ax.set_xticklabels([label for _, label in attacks], rotation=25, ha="right", fontsize=8)
    ax.set_ylabel(ylabel)
    ax.set_title(title)
    if ylim:
        ax.set_ylim(*ylim)
    ax.grid(axis="y", alpha=0.3)
    ax.legend()
    fig.tight_layout()
    fig.savefig(path, dpi=140)
    plt.close(fig)


def make_charts(rows: List[Dict], chart_dir: Path, prefix: str = "") -> Dict[str, str]:
    """Create the 4 charts from actual result rows. Returns {chart_key: filename}."""
    chart_dir.mkdir(parents=True, exist_ok=True)
    summary = summarize(rows)
    n_images = max(s["n"] for s in summary) if summary else 0
    suffix = f"(rata-rata {n_images} citra)" if n_images > 1 else "(1 citra)"
    files = {
        "psnr": f"{prefix}attack_vs_psnr.png",
        "nc": f"{prefix}attack_vs_nc.png",
        "ber": f"{prefix}attack_vs_ber.png",
        "success": f"{prefix}attack_vs_success.png",
    }
    for s in summary:  # success -> persentase
        s["success_rate"] = s["success"] / s["n"] * 100
    _grouped_bar(summary, "mean_psnr", "PSNR setelah serangan (dB)", f"Attack vs PSNR {suffix}", chart_dir / files["psnr"])
    _grouped_bar(summary, "mean_nc", "NC", f"Attack vs NC {suffix}", chart_dir / files["nc"], ylim=(0, 1.1))
    _grouped_bar(summary, "mean_ber", "BER", f"Attack vs BER {suffix}", chart_dir / files["ber"], ylim=(0, 1.0))
    _grouped_bar(summary, "success_rate", "Watermark terbaca sempurna (%)", "Attack vs tingkat keberhasilan ekstraksi", chart_dir / files["success"], ylim=(0, 115))
    return files


# ------------------------------------------------------------------- xlsx
_FONT = "Arial"
_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_OK_FILL = PatternFill("solid", fgColor="D9F2E3")
_FAIL_FILL = PatternFill("solid", fgColor="F8D7DA")
_THIN = Side(style="thin", color="B8C4D6")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

RESULT_COLUMNS = ["Image", "Method", "Attack", "Parameter", "PSNR Watermarked (dB)", "PSNR After Attack (dB)",
                  "NC", "BER", "Extracted Watermark", "Status"]


def _num(value):
    """Numbers stay numeric; infinity becomes text so Excel does not choke."""
    if isinstance(value, float) and math.isinf(value):
        return "inf"
    return value


def _style_header(ws, row: int, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=_FONT, bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER


def _autowidth(ws, minimum: int = 10, maximum: int = 42) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[letter].width = max(minimum, min(maximum, longest + 2))


def _write_results_sheet(ws, rows: List[Dict], extra: bool = False) -> None:
    columns = RESULT_COLUMNS + (["File", "SHA-256 (12)", "Message"] if extra else [])
    ws.append(columns)
    _style_header(ws, 1, len(columns))
    for r in rows:
        values = [r["image"], r["method"], r["attack"], r["parameter"], _num(r["psnr_watermarked"]),
                  _num(r["psnr_after_attack"]), r["nc"], r["ber"], r["extracted_watermark"], r["status"]]
        if extra:
            values += [r["file"], r["file_sha256"][:12], r["message"]]
        ws.append(values)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=_FONT, size=10)
            cell.border = _BORDER
        row[4].number_format = row[5].number_format = "0.00"
        row[6].number_format = row[7].number_format = "0.0000"
        row[9].fill = _OK_FILL if row[9].value == "SUCCESS" else _FAIL_FILL
        row[9].alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autowidth(ws)


def write_xlsx(results: Dict, path: Path) -> Path:
    """Write the workbook: Summary, DCT Results, LSB Results, Attack Results, Dataset."""
    rows = results["rows"]
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Digital Watermarking - Robust DCT Image Protection"
    ws["A1"].font = Font(name=_FONT, size=14, bold=True)
    info = [
        ("Dibuat pada", results.get("generated_at", "-")),
        ("Teks watermark", results.get("watermark_text", "-")),
        ("Metode", ", ".join(results.get("methods", []))),
        ("Jumlah citra", len(results.get("dataset", []))),
        ("Secret key", "tidak disimpan di laporan"),
        ("Embed strength (DCT)", results.get("config", {}).get("embed_strength", "-")),
        ("Parameter serangan", ", ".join(f"{a['short']}" for a in results.get("attacks", [])[1:])),
    ]
    for i, (label, value) in enumerate(info, start=3):
        ws.cell(row=i, column=1, value=label).font = Font(name=_FONT, bold=True)
        ws.cell(row=i, column=2, value=value).font = Font(name=_FONT)
    start = 3 + len(info) + 1
    headers = ["Method", "Attack", "Parameter", "Images", "Mean PSNR After Attack (dB)", "Mean NC", "Mean BER", "Success"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=start, column=col, value=h)
    _style_header(ws, start, len(headers))
    for offset, s in enumerate(summarize(rows), start=1):
        values = [s["method"], s["attack"], s["parameter"], s["n"], _num(s["mean_psnr"]), s["mean_nc"], s["mean_ber"], f"{s['success']}/{s['n']}"]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=start + offset, column=col, value=value)
            cell.font = Font(name=_FONT, size=10)
            cell.border = _BORDER
        ws.cell(row=start + offset, column=5).number_format = "0.00"
        ws.cell(row=start + offset, column=6).number_format = ws.cell(row=start + offset, column=7).number_format = "0.0000"
    note_row = start + len(summarize(rows)) + 2
    notes = [
        "Catatan: semua angka dihitung dari citra dan file serangan aktual (tidak ada nilai manual).",
        "PSNR Watermarked = citra asli vs citra ber-watermark. PSNR After Attack = citra asli vs citra setelah serangan.",
        "NC dan BER dihitung antara bit payload asli dan bit hasil ekstraksi (header + body).",
        "Status SUCCESS = teks watermark terbaca persis sama dan lolos CRC.",
    ]
    for i, note in enumerate(notes):
        ws.cell(row=note_row + i, column=1, value=note).font = Font(name=_FONT, italic=True, size=9)
    _autowidth(ws, maximum=34)
    ws.column_dimensions["A"].width = 24

    _write_results_sheet(wb.create_sheet("DCT Results"), [r for r in rows if r["method_key"] == "dct"])
    _write_results_sheet(wb.create_sheet("LSB Results"), [r for r in rows if r["method_key"] == "lsb"])
    _write_results_sheet(wb.create_sheet("Attack Results"), rows, extra=True)

    ds = wb.create_sheet("Dataset")
    ds.append(["Image", "File", "Width", "Height", "Size (KB)", "SHA-256 (12)", "Note"])
    _style_header(ds, 1, 7)
    for item in results.get("dataset", []):
        ds.append([item["image"], item["file"], item["width"], item["height"], item["size_kb"], item["sha256"], item["note"]])
    for row in ds.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=_FONT, size=10)
            cell.border = _BORDER
    ds.freeze_panes = "A2"
    _autowidth(ds)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
